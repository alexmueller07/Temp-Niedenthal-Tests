"""Load the CFD norming workbook: human ratings and physical measurements.

The workbook has a decorative preamble, so the header row is found rather than
assumed. Ratings are of the NEUTRAL photograph — that matters when interpreting
them: 'Happy' here means how pleasant a person looks at rest, not how big their
smile is.
"""

from __future__ import annotations

from pathlib import Path

XLSX = Path(r"C:\Users\amuel\Downloads\cfd\CFD Version 3.0"
            r"\CFD 3.0 Norming Data and Codebook.xlsx")

SHEETS = ["CFD U.S. Norming Data", "CFD-MR U.S. Norming Data", "CFD-I U.S. Norming Data"]

# Physical measurements plausibly related to how far a mouth corner can travel,
# plus the image-level luminance/colour terms that bear on whether a given
# displacement is *visible* (the beard / low-contrast question).
NUMERIC_COVARIATES = [
    "LipThickness", "LipFullness", "FaceWidthMouth", "FaceWidthCheeks",
    "FaceWidthBZ", "FaceLength", "UpperFaceLength2", "MidfaceLength",
    "ChinLength", "BottomLipChin", "CheeksAvg", "MidcheekChinR", "MidcheekChinL",
    "CheekboneProminence", "CheekboneHeight", "FaceRoundness", "fWHR2",
    "NoseWidth", "NoseLength", "EyeDistance", "PupilLipAvg", "PupilLipAsymmetry",
    "LuminanceMedian", "FaceColorRed", "FaceColorGreen", "FaceColorBlue",
    "AgeSelf", "AgeRated",
]

# Perceptual ratings of the neutral face. 'Happy' at rest is the closest thing
# the database has to resting mouth-corner curvature as judged by people.
RATING_COVARIATES = [
    "Happy", "Attractive", "Dominant", "Trustworthy", "Warm", "Babyfaced",
    "Masculine", "Feminine", "Prototypic", "Unusual",
]


def load(xlsx: Path = XLSX) -> dict[str, dict[str, float]]:
    """Return {model_id: {variable: value}}, merged across the three sheets."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    out: dict[str, dict[str, float]] = {}

    for sheet in SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))

        hdr_i = next(
            (i for i, r in enumerate(rows)
             if r and str(r[0]).strip() == "Model"),
            None,
        )
        if hdr_i is None:
            continue
        header = [str(v).strip() if v is not None else "" for v in rows[hdr_i]]
        col = {name: j for j, name in enumerate(header) if name}

        for r in rows[hdr_i + 1:]:
            if not r or r[0] is None:
                continue
            model = str(r[0]).strip()
            if not model or model in out:
                continue  # first row per model; later rows are alternate shots
            rec: dict[str, float] = {}
            for name in NUMERIC_COVARIATES + RATING_COVARIATES:
                j = col.get(name)
                if j is None or j >= len(r):
                    continue
                try:
                    v = float(r[j])
                except (TypeError, ValueError):
                    continue
                rec[name] = v
            if rec:
                out[model] = rec

    return out


if __name__ == "__main__":
    d = load()
    print(f"{len(d)} models with norming data")
    if d:
        k = next(iter(d))
        print(f"example {k}: {len(d[k])} variables")
        for name in ("LipThickness", "FaceWidthMouth", "Happy", "LuminanceMedian"):
            print(f"  {name} = {d[k].get(name)}")
